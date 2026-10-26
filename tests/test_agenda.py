"""Extractores de agenda: regla C, texto interno, layouts y CSV público."""

from __future__ import annotations

import csv
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from datetime import date, time
from io import StringIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color

from sesiones.agenda import (
    COLUMNAS_CERRADAS,
    COLUMNAS_PUBLICAS,
    COLUMNAS_SLOTS_CERRADAS,
    COLUMNAS_SLOTS_PUBLICAS,
    ejecutar,
    ejecutar_clasificacion,
    exportar_slots_publicos,
    extraer_horizontal,
    extraer_vertical,
    filas_revision,
)
from sesiones.bloques import BloqueHorizontal, BloqueVertical
from sesiones.cli import main

VERDE = "FF6AA84F"
ROJO = "FFFF0000"

FOLIO_IU = "CEXXIU001"
FOLIO_IF = "CEXXIF073"


def relleno(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def libro_vertical() -> Workbook:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Semana 1"
    hoja["B6"] = "THERAPIST_1 SITE_A"
    hoja["C6"] = "THERAPIST_2 SITE_A NORTH"
    hoja["A8"] = time(10, 0)
    hoja["A9"] = time(11, 0)
    hoja["A10"] = time(12, 0)
    hoja["B8"] = "USER_1 10 a 11"
    hoja["B8"].fill = relleno(VERDE)
    hoja["B9"] = "USER_1 y USER_2"
    hoja["B9"].fill = relleno(VERDE)
    hoja["C8"] = "sin marca"
    hoja["C9"] = "USER_3"
    hoja["C9"].fill = relleno(ROJO)
    hoja["B10"] = None
    hoja["B10"].fill = relleno(VERDE)
    hoja["C10"] = "supervision meeting"
    hoja["C10"].fill = relleno(VERDE)
    hoja["A17"] = time(9, 0)
    hoja["A19"] = time(10, 0)
    hoja["B15"] = "THERAPIST_1 SITE_A"
    hoja["C15"] = "THERAPIST_2 SITE_A NORTH"
    hoja["B17"] = f"{FOLIO_IU} USER_4"
    hoja["B17"].fill = relleno(VERDE)
    hoja["C19"] = "USER_5"
    hoja["C19"].fill = relleno(VERDE)
    hoja["B19"] = "Tema"
    hoja["B19"].fill = PatternFill(fill_type="solid", fgColor=Color(theme=4, tint=0.4))
    return libro


BLOQUES_V = (
    BloqueVertical("Semana 1", date(2025, 10, 1), 6, 8, 11),
    BloqueVertical("Semana 1", date(2025, 10, 2), 15, 17, 20),
)
SEDES = ("SITE_A NORTH", "SITE_A", "SITE_B")


class Vertical(unittest.TestCase):
    def test_regla_c_y_texto_interno(self):
        sesiones = extraer_vertical(libro_vertical(), SEDES, BLOQUES_V)
        self.assertEqual(len(sesiones), 4)
        self.assertTrue(all(s.es_sesion for s in sesiones))
        self.assertEqual(sesiones[0].hora, time(10, 0))
        self.assertIsNone(sesiones[0].folio)
        self.assertEqual(sesiones[2].folio, FOLIO_IU)
        self.assertEqual(sesiones[2].iu_if, "IU")
        self.assertEqual(sesiones[2].layout, "vertical")
        self.assertEqual(sesiones[3].sede, "SITE_A")

    def test_csv_publico_sin_nombres(self):
        resultado = ejecutar(libro_vertical(), SEDES, BLOQUES_V)
        filas = resultado.filas
        self.assertEqual(list(filas[0]), COLUMNAS_PUBLICAS)
        self.assertEqual([f["contacto_id"] for f in filas], [1, 2, 3, 4])
        self.assertEqual(filas[0]["terapeuta_alias"], filas[1]["terapeuta_alias"])
        self.assertEqual(filas[2]["folio"], FOLIO_IU)
        self.assertEqual(filas[2]["iu_if"], "IU")
        texto = " ".join(str(v) for f in filas for v in f.values())
        for secreto in ("USER_", "THERAPIST_", "SITE_", "meeting"):
            self.assertNotIn(secreto, texto)
        self.assertIn("2025-10-01 2", resultado.informe)

    def test_sede_nueva_no_se_rechaza(self):
        libro = libro_vertical()
        libro["Semana 1"]["D6"] = "THERAPIST_3"
        libro["Semana 1"]["D8"] = "contacto anonimo"
        libro["Semana 1"]["D8"].fill = relleno(VERDE)
        sesiones = extraer_vertical(libro, SEDES, BLOQUES_V)
        self.assertEqual(len(sesiones), 5)
        nueva = next(s for s in sesiones if s.terapeuta == "THERAPIST_3")
        self.assertIsNone(nueva.sede)
        filas = ejecutar(libro, SEDES, BLOQUES_V).filas
        texto = " ".join(str(v) for f in filas for v in f.values())
        self.assertNotIn("contacto anonimo", texto)
        self.assertNotIn("THERAPIST_3", texto)


class Horizontal(unittest.TestCase):
    def test_dias_lado_a_lado(self):
        libro = Workbook()
        hoja = libro.active
        hoja.title = "FEBRERO"
        hoja["A9"] = "HORARIOS"
        hoja["B9"] = "THERAPIST_1 SITE_A"
        hoja["C9"] = "THERAPIST_2 SITE_A"
        hoja["O9"] = "HORARIOS"
        hoja["P9"] = "THERAPIST_1 SITE_A"
        hoja["A11"] = time(10, 0)
        hoja["A12"] = time(18, 0)
        hoja["O11"] = time(10, 0)
        hoja["B11"] = f"{FOLIO_IF} uno"
        hoja["B11"].fill = relleno(VERDE)
        hoja["C12"] = "institutional workshop"
        hoja["C12"].fill = relleno(VERDE)
        hoja["P11"] = "dos"
        hoja["P11"].fill = relleno(VERDE)
        bloques = (
            BloqueHorizontal("FEBRERO", date(2026, 2, 3), 1, 15, 9, 11, 13),
            BloqueHorizontal("FEBRERO", date(2026, 2, 4), 15, 29, 9, 11, 13),
        )
        sesiones = extraer_horizontal(libro, ("SITE_A",), bloques)
        self.assertEqual(len(sesiones), 2)
        self.assertEqual(sesiones[0].fecha, date(2026, 2, 3))
        self.assertEqual(sesiones[0].folio, FOLIO_IF)
        self.assertEqual(sesiones[0].iu_if, "IF")
        self.assertEqual(sesiones[0].hora, time(10, 0))
        self.assertEqual(sesiones[1].fecha, date(2026, 2, 4))
        self.assertEqual(sesiones[1].layout, "horizontal")


class ClasificacionAgenda(unittest.TestCase):
    def test_csv_publico_clasificacion_sin_nombres(self):
        resultado = ejecutar_clasificacion(
            libro_vertical(),
            SEDES,
            BLOQUES_V,
            "cartelera_ejemplo.xlsx",
        )
        filas = resultado.filas_publicas
        self.assertEqual(list(filas[0].keys()), COLUMNAS_SLOTS_PUBLICAS)
        texto = " ".join(str(v) for f in filas for v in f.values())
        for secreto in ("USER_", "THERAPIST_", "SITE_", "meeting", "CEXX"):
            self.assertNotIn(secreto, texto)
        self.assertNotIn("folio", filas[0])
        self.assertNotIn("iu_if", filas[0])
        estados = {f["estado"] for f in filas}
        self.assertIn("asiste", estados)

    def test_hora_invalida_no_aborta_y_va_a_revision(self):
        libro = libro_vertical()
        libro["Semana 1"]["A8"] = "no es hora"
        libro["Semana 1"]["B8"] = "USER_6 con hora mala"
        libro["Semana 1"]["B8"].fill = relleno(VERDE)
        resultado = ejecutar_clasificacion(libro, SEDES, BLOQUES_V[:1], "cartelera_ejemplo.xlsx")
        self.assertEqual(resultado.filas_hora_invalida[0]["motivo"], "hora_invalida")
        self.assertEqual(resultado.filas_hora_invalida[0]["celda"], "B8")
        revision = filas_revision(resultado.filas_internas, resultado.filas_hora_invalida)
        self.assertIn(resultado.filas_hora_invalida[0], revision)

    def test_columnas_cerradas_como_oct_nov(self):
        filas = exportar_slots_publicos(
            ejecutar_clasificacion(
                libro_vertical(),
                SEDES,
                BLOQUES_V,
                "cartelera_ejemplo.xlsx",
            ).slots,
            cerradas=True,
        )
        self.assertEqual(list(filas[0].keys()), COLUMNAS_SLOTS_CERRADAS)
        base = {k for k in COLUMNAS_SLOTS_CERRADAS if k in COLUMNAS_CERRADAS}
        self.assertEqual(base, set(COLUMNAS_CERRADAS))


class LineaDeComandos(unittest.TestCase):
    def test_escribe_el_csv_del_mes(self):
        from sesiones.bloques import MESES

        with tempfile.TemporaryDirectory() as tmp:
            directorio = Path(tmp)
            libro = libro_vertical()
            ruta_libro = directorio / "cartelera_ejemplo.xlsx"
            libro.save(ruta_libro)
            (directorio / "sedes.txt").write_text("SITE_B\nSITE_A\nSITE_A NORTH\n", encoding="utf-8")
            from sesiones.agenda import procesar

            resultado = procesar(
                ruta_libro,
                directorio / "sedes.txt",
                BLOQUES_V,
                "vertical",
            )
            salida = directorio / "out" / "sesiones.csv"
            from sesiones.agenda import escribir_csv

            escribir_csv(salida, resultado.filas)
            with salida.open(encoding="utf-8", newline="") as archivo:
                filas = list(csv.DictReader(archivo))
            self.assertEqual(list(filas[0]), COLUMNAS_PUBLICAS)
            self.assertEqual(len(filas), 4)
            self.assertNotIn("USER_", salida.read_text(encoding="utf-8"))
            self.assertIn("oct2025", MESES)

    def test_mes_desconocido(self):
        with redirect_stdout(StringIO()), redirect_stderr(StringIO()) as err:
            codigo = main(
                [
                    "agenda",
                    "--workbook",
                    "ejemplo.xlsx",
                    "--sites",
                    "sedes.txt",
                    "--month",
                    "julio2099",
                    "--output",
                    "salida.csv",
                ]
            )
        self.assertEqual(codigo, 1)
        self.assertIn("julio2099", err.getvalue())


if __name__ == "__main__":
    unittest.main()
