"""Reglas del modelo: sede, folio, texto excluido, color."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color

from sesiones.reglas import (
    ErrorDeModelo,
    cargar_sedes,
    clasificar_celda,
    clasificar_slot,
    es_texto_excluido,
    extraer_folio,
    extraer_iu_if,
    normalizar_sede,
    partir_encabezado,
)

FOLIO_IF = "CEXXIF054"
FOLIO_IU = "CEXXIU001"
FOLIO_IC = "CEXXIC004"


class Sede(unittest.TestCase):
    def test_variantes_del_modelo(self):
        self.assertEqual(normalizar_sede("SITE_A_TYPO1"), "SITE_A")
        self.assertEqual(normalizar_sede("SITE_A_TYPO2"), "SITE_A")
        self.assertEqual(normalizar_sede("SITE_B_EXTENDED"), "SITE_B")
        self.assertEqual(normalizar_sede("branch label - site_a_typo2"), "SITE_A")

    def test_partir_la_sede_mas_larga(self):
        sedes = ("SITE_A NORTH", "SITE_A", "SITE_B")
        self.assertEqual(
            partir_encabezado("THERAPIST_2 SITE_A NORTH", sedes),
            ("THERAPIST_2", "SITE_A"),
        )
        self.assertEqual(partir_encabezado("THERAPIST_1 SITE_A", sedes), ("THERAPIST_1", "SITE_A"))


class Folio(unittest.TestCase):
    def test_extrae_iu_e_if(self):
        self.assertEqual(extraer_folio(FOLIO_IF), FOLIO_IF)
        self.assertEqual(extraer_folio(f"cita {FOLIO_IU} 10am"), FOLIO_IU)
        self.assertEqual(extraer_folio("IU-12"), "IU12")
        self.assertEqual(extraer_iu_if(FOLIO_IF), "IF")
        self.assertEqual(extraer_iu_if(FOLIO_IU), "IU")
        self.assertEqual(extraer_iu_if(FOLIO_IC), None)

    def test_no_rellena_un_nombre(self):
        self.assertIsNone(extraer_folio("USER_1 10:00-10:45"))
        self.assertIsNone(extraer_folio("NO APLICA"))
        self.assertIsNone(extraer_folio(None))


class TextoExcluido(unittest.TestCase):
    def test_trabajo_interno(self):
        self.assertTrue(es_texto_excluido("team meeting"))
        self.assertTrue(es_texto_excluido("prevention workshop"))
        self.assertTrue(es_texto_excluido("lunch break"))
        self.assertFalse(es_texto_excluido("USER_1 10 a 11"))
        self.assertFalse(es_texto_excluido(None))


class Clasificacion(unittest.TestCase):
    def test_vacio(self):
        self.assertEqual(clasificar_slot("", None).estado, "vacio")

    def test_asiste_verde_cartelera(self):
        c = clasificar_slot("USER_1 10 a 11", "FF6AA84F")
        self.assertEqual(c.estado, "asiste")

    def test_asiste_verde_brillante(self):
        c = clasificar_slot("USER_2", "FF00FF00")
        self.assertEqual(c.estado, "asiste")

    def test_asiste_verde_tema_accent4(self):
        c = clasificar_slot("USER_3", "FF34A853")
        self.assertEqual(c.estado, "asiste")

    def test_reagendo_amarillo(self):
        c = clasificar_slot("USER_4 Apellido", "FFFFFF00")
        self.assertEqual(c.estado, "reagendo")
        self.assertNotIn("revisar_texto", c.flags)

    def test_inasistencia_rojo(self):
        self.assertEqual(clasificar_slot("USER_5", "FFFF0000").estado, "inasistencia")

    def test_egreso_azul(self):
        self.assertEqual(clasificar_slot("Alta", "FF1155CC").estado, "egreso")
        self.assertEqual(clasificar_slot("Alta", "FF0000FF").estado, "egreso")

    def test_baja_naranja(self):
        self.assertEqual(clasificar_slot("Baja", "FFED7D31").estado, "baja")

    def test_ff9900_baja_vertical(self):
        self.assertEqual(clasificar_slot("Motivo", "FFFF9900", layout="vertical").estado, "baja")

    def test_ff9900_ingreso_horizontal(self):
        self.assertEqual(clasificar_slot("Nuevo", "FFFF9900", layout="horizontal").estado, "ingreso")

    def test_descanso_no_sesion(self):
        c = clasificar_slot("scheduled break", "FFFFFF00", layout="horizontal")
        self.assertEqual(c.estado, "no_sesion")
        self.assertEqual(c.motivo, "texto_excluido")

    def test_no_sesion_magenta_comida(self):
        self.assertEqual(clasificar_slot("Comida", "FFFF00FF").estado, "no_sesion")

    def test_texto_gana_sobre_color(self):
        c = clasificar_slot("MEETING SITE_A", "FFFF0000")
        self.assertEqual(c.estado, "no_sesion")
        self.assertEqual(c.motivo, "texto_excluido")

    def test_fondo_sin_estado(self):
        c = clasificar_slot("USER_1 10 a 11", "FFD9EAD3")
        self.assertEqual(c.estado, "sin_clasificar")
        self.assertEqual(c.motivo, "texto_sin_color")

    def test_fondo_estado_por_texto(self):
        c = clasificar_slot("Usuario (alta)", "FFD9EAD3")
        self.assertEqual(c.estado, "egreso")
        self.assertIn("estado_por_texto", c.flags)

    def test_ingreso_dorado(self):
        c = clasificar_slot(f"USER_6 {FOLIO_IU}", "FFFFC000")
        self.assertEqual(c.estado, "ingreso")

    def test_asiste_93c47d(self):
        self.assertEqual(clasificar_slot("USER_7", "FF93C47D").estado, "asiste")

    def test_baja_ff6d01_aproximado(self):
        c = clasificar_slot("USER_8", "FFFF6D01")
        self.assertEqual(c.estado, "baja")
        self.assertIn("color_aproximado", c.flags)

    def test_apoyo_en_curso_revisar_texto(self):
        c = clasificar_slot("Apoyo en Curso de Verano", "FF6AA84F")
        self.assertEqual(c.estado, "asiste")
        self.assertIn("revisar_texto", c.flags)

    def test_color_desconocido(self):
        c = clasificar_slot("algo", "FFFF8000")
        self.assertEqual(c.estado, "sin_clasificar")
        self.assertEqual(c.motivo, "color_desconocido")

    def test_fondo_rosa_sin_palabra_clave(self):
        c = clasificar_slot("algo", "FFE6B8AF")
        self.assertEqual(c.estado, "sin_clasificar")
        self.assertEqual(c.motivo, "texto_sin_color")

    def test_conflicto_color_texto(self):
        c = clasificar_slot("USER_9 (alta)", "FFFF0000")
        self.assertEqual(c.estado, "inasistencia")
        self.assertIn("conflicto_color_texto", c.flags)

    def test_reagendo_revisar_texto(self):
        c = clasificar_slot("cons grupal", "FFFFFF00")
        self.assertEqual(c.estado, "reagendo")
        self.assertIn("revisar_texto", c.flags)

    def test_relleno_no_resuelto(self):
        c = clasificar_slot("Tema", None, resuelto=False)
        self.assertEqual(c.estado, "sin_clasificar")
        self.assertEqual(c.motivo, "color_desconocido")

    def test_theme_accent4_resuelve_asiste(self):
        libro = Workbook()
        libro._sesiones_paleta_tema = tuple([None] * 7 + ["34A853"] + [None] * 4)
        hoja = libro.active
        hoja["A1"] = "USER_3"
        hoja["A1"].fill = PatternFill(fill_type="solid", fgColor=Color(theme=7, tint=0.0))
        c = clasificar_celda(hoja["A1"], libro)
        self.assertEqual(c.estado, "asiste")
        libro.close()


class Catalogo(unittest.TestCase):
    def test_variante_fuera_del_catalogo(self):
        with tempfile.TemporaryDirectory() as tmp:
            sedes = Path(tmp) / "sedes.txt"
            sedes.write_text("SITE_A\n", encoding="utf-8")
            variantes = Path(tmp) / "var.txt"
            variantes.write_text("SITE_B=SITE_A\n", encoding="utf-8")
            with self.assertRaises(ErrorDeModelo):
                cargar_sedes(sedes, variantes)
            sedes.write_text("\n", encoding="utf-8")
            with self.assertRaises(ErrorDeModelo):
                cargar_sedes(sedes)


if __name__ == "__main__":
    unittest.main()
