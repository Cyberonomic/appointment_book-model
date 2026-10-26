"""Lectores de expediente y cruce solo por folio."""

from __future__ import annotations

import unittest
from datetime import date, datetime, time

from openpyxl import Workbook

from sesiones.agenda import SesionAgenda
from sesiones.expedientes import (
    COL_2025,
    COL_2026,
    leer_expediente_2025,
    leer_expediente_2026,
    mapear_estatus,
    mapear_tipo,
)
from sesiones.exportar import (
    cruzar_por_folio,
    exportar_eventos_2025_publicos,
    exportar_personas_2026_publicas,
    unificar_personas,
)

FOLIO_2025 = "CEXXIU061"
FOLIO_2026 = "CEXXIU001"
HOJA_2025 = "Sheet1"
HOJA_2026 = "Data"


def libro_2025() -> Workbook:
    libro = Workbook()
    hoja = libro.active
    hoja.title = HOJA_2025
    fila = (
        (COL_2025["terapeuta"], "THERAPIST_1"),
        (COL_2025["sede"], "SITE_B EXTENDED"),
        (COL_2025["fecha"], datetime(2025, 12, 2)),
        (COL_2025["tipo"], "Acompañamiento psicológico individual (usuario)"),
        (COL_2025["folio"], FOLIO_2025),
        (COL_2025["estatus"], "Sesión de acompañamiento psicológico"),
        (COL_2025["sexo"], "Mujer"),
        (COL_2025["rango"], "18 a 29"),
        (COL_2025["sustancia"], "CANNABIS"),
        (COL_2025["numero_sesion"], 3),
    )
    for col, valor in fila:
        hoja.cell(2, col, valor)
    hoja.cell(3, COL_2025["tipo"], "Consejeria")
    hoja.cell(3, COL_2025["estatus"], "Alta por termino de proceso")
    hoja.cell(3, COL_2025["fecha"], datetime(2025, 12, 3))
    hoja.cell(3, COL_2025["folio"], "texto que no es folio")
    return libro


def libro_2026() -> Workbook:
    libro = Workbook()
    hoja = libro.active
    hoja.title = HOJA_2026
    hoja.cell(2, COL_2026["folio"], FOLIO_2026)
    hoja.cell(2, COL_2026["terapeuta"], "THERAPIST_1")
    hoja.cell(2, COL_2026["sede"], "branch label - site_a_typo2")
    hoja.cell(2, COL_2026["edad"], 35)
    hoja.cell(2, COL_2026["genero"], "MUJER")
    hoja.cell(2, COL_2026["sustancia"], "ALCOHOL")
    hoja.cell(2, COL_2026["fecha_primera"], datetime(2026, 4, 15))
    hoja.cell(2, COL_2026["fecha_ultima"], datetime(2026, 5, 19))
    hoja.cell(2, COL_2026["sesiones"][4], 1)
    hoja.cell(2, COL_2026["sesiones"][5], 2)
    hoja.cell(2, COL_2026["total"], 3)
    hoja.cell(2, COL_2026["estatus_tratamiento"], "EN TRATAMIENTO")
    hoja.cell(2, COL_2026["motivo_egreso"], None)
    hoja.cell(2, 51, "0000000000")
    hoja.cell(3, COL_2026["folio"], None)
    hoja.cell(3, COL_2026["terapeuta"], "sin folio")
    return libro


class Mapeos(unittest.TestCase):
    def test_tipo_y_estatus(self):
        self.assertEqual(mapear_tipo("Acompañamiento psicológico individual (familiar)"), "familiar")
        self.assertEqual(mapear_tipo("Consejeria"), "consejería")
        self.assertEqual(mapear_estatus("Alta por inicio de proceso"), "alta por inicio")
        self.assertEqual(mapear_estatus("Asesoria general"), "asesoría")
        self.assertTrue(leer_expediente_2025(libro_2025())[0].es_sesion)
        self.assertFalse(leer_expediente_2025(libro_2025())[1].es_sesion)


class Lectura(unittest.TestCase):
    def test_2025_folio_nulo_se_queda_nulo(self):
        eventos = leer_expediente_2025(libro_2025())
        self.assertEqual(len(eventos), 2)
        self.assertEqual(eventos[0].folio, FOLIO_2025)
        self.assertEqual(eventos[0].sede, "SITE_B")
        self.assertEqual(eventos[0].tipo, "usuario")
        self.assertIsNone(eventos[1].folio)
        self.assertEqual(eventos[1].estatus, "alta por término")
        publico = exportar_eventos_2025_publicos(eventos)
        texto = " ".join(str(v) for f in publico for v in f.values())
        self.assertNotIn("THERAPIST_1", texto)
        self.assertNotIn("texto que no es folio", texto)

    def test_2026_sin_telefono_ni_nombre(self):
        personas = leer_expediente_2026(libro_2026())
        self.assertEqual(len(personas), 1)
        self.assertEqual(personas[0].folio, FOLIO_2026)
        self.assertEqual(personas[0].iu_if, "IU")
        self.assertEqual(personas[0].total_sesiones, 3)
        self.assertEqual(personas[0].sesiones_mes[5], 2)
        self.assertEqual(personas[0].sede, "SITE_A")
        publico = exportar_personas_2026_publicas(personas)
        texto = " ".join(str(v) for f in publico for v in f.values())
        self.assertNotIn("THERAPIST_1", texto)
        self.assertNotIn("0000000000", texto)
        self.assertNotIn("sin folio", texto)


class Cruce(unittest.TestCase):
    def test_solo_por_folio(self):
        eventos = leer_expediente_2025(libro_2025())
        personas = leer_expediente_2026(libro_2026())
        unificadas = unificar_personas(eventos, personas)
        folios = {p.folio: p.origen for p in unificadas}
        self.assertEqual(folios[FOLIO_2026], "2026")
        self.assertEqual(folios[FOLIO_2025], "2025")

        sesion_con = SesionAgenda(
            1,
            date(2026, 1, 10),
            "THERAPIST_1",
            "SITE_A",
            time(10, 0),
            "x",
            FOLIO_2026,
            "IU",
            "vertical",
            True,
        )
        sesion_sin = SesionAgenda(
            2,
            date(2025, 10, 1),
            "THERAPIST_1",
            "SITE_A",
            time(11, 0),
            "y",
            None,
            None,
            "vertical",
            True,
        )
        cruce = cruzar_por_folio([sesion_con, sesion_sin], unificadas)
        self.assertIsNotNone(cruce[0][1])
        self.assertEqual(cruce[0][1].folio, FOLIO_2026)
        self.assertIsNone(cruce[1][1])


if __name__ == "__main__":
    unittest.main()
