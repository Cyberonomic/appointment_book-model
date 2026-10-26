"""Extractores de agenda vertical (oct–ene) y horizontal (feb+)."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path

from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook

from sesiones.bloques import BloqueHorizontal, BloqueVertical
from sesiones.reglas import (
    ErrorDeModelo,
    clasificar_celda,
    codigo,
    color_rgb,
    es_asiste,
    es_sesion,
    extraer_folio,
    extraer_iu_if,
    partir_encabezado,
)

ErrorDeSesiones = ErrorDeModelo

COLUMNAS_CERRADAS = ["contacto_id", "fecha", "terapeuta_alias", "sede_alias", "hora"]
COLUMNAS_PUBLICAS = [
    "contacto_id",
    "fecha",
    "terapeuta_alias",
    "sede_alias",
    "hora",
    "folio",
    "iu_if",
    "layout",
]

COLUMNAS_SLOTS_INTERNAS = [
    "archivo",
    "hoja",
    "celda",
    "layout",
    "fecha",
    "dia",
    "hora",
    "terapeuta",
    "sede",
    "texto_original",
    "color_hex",
    "estado",
    "motivo",
    "flags",
    "folio",
    "tipo_folio",
]

COLUMNAS_SLOTS_PUBLICAS = [
    "contacto_id",
    "fecha",
    "terapeuta_alias",
    "sede_alias",
    "hora",
    "layout",
    "estado",
    "motivo",
    "flags",
    "hoja",
    "celda",
    "archivo",
]

COLUMNAS_SLOTS_CERRADAS = [
    "contacto_id",
    "fecha",
    "terapeuta_alias",
    "sede_alias",
    "hora",
    "estado",
    "motivo",
    "flags",
    "hoja",
    "celda",
    "archivo",
]

_DIAS_ES = (
    "lunes",
    "martes",
    "miercoles",
    "jueves",
    "viernes",
    "sabado",
    "domingo",
)


@dataclass(frozen=True)
class SesionAgenda:
    contacto_id: int
    fecha: date
    terapeuta: str
    sede: str | None
    hora: time
    texto: str
    folio: str | None
    iu_if: str | None
    layout: str
    es_sesion: bool


@dataclass(frozen=True)
class SlotClasificado:
    archivo: str
    hoja: str
    celda: str
    layout: str
    fecha: date
    dia: str
    hora: time
    terapeuta: str
    sede: str | None
    texto_original: str
    color_hex: str
    estado: str
    motivo: str
    flags: str
    folio: str | None
    tipo_folio: str | None


@dataclass
class Resultado:
    sesiones: list[SesionAgenda]
    filas: list[dict[str, str | int]]
    dias: tuple[date, ...]

    @property
    def informe(self) -> str:
        por_dia = Counter(fila["fecha"] for fila in self.filas)
        terapeutas = {fila["terapeuta_alias"] for fila in self.filas}
        sedes = {fila["sede_alias"] for fila in self.filas}
        folios = sum(1 for fila in self.filas if fila["folio"])
        lineas = [
            f"{len(self.filas)} sesiones efectuadas",
            f"{self.dias[0].isoformat()} a {self.dias[-1].isoformat()}",
            f"{len(terapeutas)} terapeutas, {len(sedes)} sedes, {folios} con folio",
        ]
        lineas.extend(f"{dia.isoformat()} {por_dia.get(dia.isoformat(), 0)}" for dia in self.dias)
        return "\n".join(lineas)


@dataclass
class ResultadoClasificado:
    slots: list[SlotClasificado]
    filas_internas: list[dict[str, str]]
    filas_publicas: list[dict[str, str | int]]
    filas_hora_invalida: list[dict[str, str]]
    dias: tuple[date, ...]
    archivo: str

    def informe_clasificacion(self) -> str:
        por_estado = Counter(s.estado for s in self.slots)
        sin_motivo = Counter(s.motivo for s in self.slots if s.estado == "sin_clasificar")
        lineas = [f"Clasificación: {self.archivo}", f"{len(self.slots)} slots"]
        if self.filas_hora_invalida:
            lineas.append(f"  hora_invalida (solo revision): {len(self.filas_hora_invalida)}")
        for estado in sorted(por_estado):
            lineas.append(f"  {estado}: {por_estado[estado]}")
        if sin_motivo:
            lineas.append("  sin_clasificar por motivo:")
            for motivo, cuenta in sorted(sin_motivo.items()):
                lineas.append(f"    {motivo or '(vacío)'}: {cuenta}")
        return "\n".join(lineas)


def _hora(valor) -> time | None:
    if isinstance(valor, datetime):
        return valor.time()
    if isinstance(valor, time):
        return valor
    return None


def _leer_slot(
    hoja,
    fila: int,
    columna: int,
    col_hora: int,
    fecha: date,
    terapeuta: str,
    sede: str | None,
    layout: str,
    excluir: bool,
    donde: str,
    errores: list[str],
) -> SesionAgenda | None:
    celda = hoja.cell(row=fila, column=columna)
    if excluir:
        if not es_sesion(celda):
            return None
    elif not es_asiste(celda):
        return None
    if not terapeuta:
        return None
    hora = _hora(hoja.cell(row=fila, column=col_hora).value)
    if hora is None:
        errores.append(f"{donde} fila {fila}: el horario no es una hora")
        return None
    texto = str(celda.value)
    folio = extraer_folio(texto)
    return SesionAgenda(
        contacto_id=0,
        fecha=fecha,
        terapeuta=terapeuta,
        sede=sede,
        hora=hora,
        texto=texto,
        folio=folio,
        iu_if=extraer_iu_if(folio),
        layout=layout,
        es_sesion=True,
    )


def _dia_semana(fecha: date) -> str:
    return _DIAS_ES[fecha.weekday()]


def _flags_str(flags: tuple[str, ...]) -> str:
    return ",".join(flags)


def _fila_hora_invalida(
    archivo: str,
    hoja_nombre: str,
    celda: str,
    texto: str,
    layout: str,
    fecha: date,
    terapeuta: str,
    sede: str | None,
    color_hex: str,
) -> dict[str, str]:
    return {
        "archivo": archivo,
        "hoja": hoja_nombre,
        "celda": celda,
        "layout": layout,
        "fecha": fecha.isoformat(),
        "dia": _dia_semana(fecha),
        "hora": "",
        "terapeuta": terapeuta,
        "sede": sede or "",
        "texto_original": texto,
        "color_hex": color_hex,
        "estado": "",
        "motivo": "hora_invalida",
        "flags": "",
        "folio": "",
        "tipo_folio": "",
    }


def _leer_slot_clasificado(
    hoja,
    fila: int,
    columna: int,
    col_hora: int,
    fecha: date,
    terapeuta: str,
    sede: str | None,
    layout: str,
    archivo: str,
    hoja_nombre: str,
    incluir_vacios: bool,
    hora_invalidas: list[dict[str, str]],
) -> SlotClasificado | None:
    celda = hoja.cell(row=fila, column=columna)
    libro = hoja.parent
    clas = clasificar_celda(celda, libro, layout)
    if clas.estado == "vacio" and not incluir_vacios:
        return None
    if not terapeuta and clas.estado != "vacio":
        return None
    texto = "" if celda.value is None else str(celda.value)
    rgb = color_rgb(celda, libro)
    ref = f"{get_column_letter(columna)}{fila}"
    hora = _hora(hoja.cell(row=fila, column=col_hora).value)
    if hora is None and clas.estado != "vacio":
        hora_invalidas.append(
            _fila_hora_invalida(
                archivo,
                hoja_nombre,
                ref,
                texto,
                layout,
                fecha,
                terapeuta,
                sede,
                rgb or "",
            )
        )
        return None
    folio = extraer_folio(texto) if texto.strip() else None
    return SlotClasificado(
        archivo=archivo,
        hoja=hoja_nombre,
        celda=ref,
        layout=layout,
        fecha=fecha,
        dia=_dia_semana(fecha),
        hora=hora or time(0, 0),
        terapeuta=terapeuta,
        sede=sede,
        texto_original=texto,
        color_hex=rgb or "",
        estado=clas.estado,
        motivo=clas.motivo,
        flags=_flags_str(clas.flags),
        folio=folio,
        tipo_folio=extraer_iu_if(folio),
    )


def _slot_a_fila_interna(slot: SlotClasificado) -> dict[str, str]:
    return {
        "archivo": slot.archivo,
        "hoja": slot.hoja,
        "celda": slot.celda,
        "layout": slot.layout,
        "fecha": slot.fecha.isoformat(),
        "dia": slot.dia,
        "hora": slot.hora.strftime("%H:%M:%S"),
        "terapeuta": slot.terapeuta,
        "sede": slot.sede or "",
        "texto_original": slot.texto_original,
        "color_hex": slot.color_hex,
        "estado": slot.estado,
        "motivo": slot.motivo,
        "flags": slot.flags,
        "folio": slot.folio or "",
        "tipo_folio": slot.tipo_folio or "",
    }


def exportar_slots_publicos(
    slots: list[SlotClasificado],
    cerradas: bool = False,
) -> list[dict[str, str | int]]:
    """Alias como sesiones públicas; sin texto, terapeuta en claro ni folio (folio solo en XLSX interno)."""
    terapeutas: dict[str, str] = {}
    sedes_cod: dict[str, str] = {}
    filas: list[dict[str, str | int]] = []
    for indice, slot in enumerate(slots, start=1):
        base: dict[str, str | int] = {
            "contacto_id": indice,
            "fecha": slot.fecha.isoformat(),
            "terapeuta_alias": codigo(terapeutas, slot.terapeuta, "T", 3),
            "sede_alias": codigo(sedes_cod, slot.sede or "", "S", 2),
            "hora": slot.hora.strftime("%H:%M:%S"),
        }
        if cerradas:
            fila = {
                **base,
                "estado": slot.estado,
                "motivo": slot.motivo,
                "flags": slot.flags,
                "hoja": slot.hoja,
                "celda": slot.celda,
                "archivo": slot.archivo,
            }
        else:
            fila = {
                **base,
                "layout": slot.layout,
                "estado": slot.estado,
                "motivo": slot.motivo,
                "flags": slot.flags,
                "hoja": slot.hoja,
                "celda": slot.celda,
                "archivo": slot.archivo,
            }
        filas.append(fila)
    return filas


def filas_revision(
    filas_internas: list[dict[str, str]],
    filas_hora_invalida: list[dict[str, str]] | None = None,
) -> list[dict[str, str]]:
    revision = [
        fila
        for fila in filas_internas
        if fila.get("estado") == "sin_clasificar" or fila.get("flags", "").strip()
    ]
    if filas_hora_invalida:
        revision.extend(filas_hora_invalida)
    return revision


def extraer_clasificado_vertical(
    libro: Workbook,
    sedes: tuple[str, ...],
    bloques: tuple[BloqueVertical, ...] | list[BloqueVertical],
    archivo: str,
    extra: dict[str, str] | None = None,
    incluir_vacios: bool = False,
) -> tuple[list[SlotClasificado], list[dict[str, str]]]:
    slots: list[SlotClasificado] = []
    hora_invalidas: list[dict[str, str]] = []
    errores: list[str] = []
    for bloque in bloques:
        if bloque.hoja not in libro.sheetnames:
            errores.append(f"no está la hoja {bloque.hoja}")
            continue
        hoja = libro[bloque.hoja]
        for columna in range(2, hoja.max_column + 1):
            crudo = hoja.cell(row=bloque.fila_encabezado, column=columna).value
            terapeuta, sede = partir_encabezado(str(crudo) if crudo is not None else None, sedes, extra)
            for fila in range(bloque.fila_inicio, bloque.fila_fin):
                slot = _leer_slot_clasificado(
                    hoja,
                    fila,
                    columna,
                    1,
                    bloque.fecha,
                    terapeuta,
                    sede,
                    "vertical",
                    archivo,
                    bloque.hoja,
                    incluir_vacios,
                    hora_invalidas,
                )
                if slot is not None:
                    slots.append(slot)
    if errores:
        raise ErrorDeModelo(errores)
    return slots, hora_invalidas


def extraer_clasificado_horizontal(
    libro: Workbook,
    sedes: tuple[str, ...],
    bloques: tuple[BloqueHorizontal, ...] | list[BloqueHorizontal],
    archivo: str,
    extra: dict[str, str] | None = None,
    incluir_vacios: bool = False,
) -> tuple[list[SlotClasificado], list[dict[str, str]]]:
    slots: list[SlotClasificado] = []
    hora_invalidas: list[dict[str, str]] = []
    errores: list[str] = []
    for bloque in bloques:
        if bloque.hoja not in libro.sheetnames:
            errores.append(f"no está la hoja {bloque.hoja}")
            continue
        hoja = libro[bloque.hoja]
        for columna in range(bloque.col_inicio + 1, bloque.col_fin):
            crudo = hoja.cell(row=bloque.fila_encabezado, column=columna).value
            terapeuta, sede = partir_encabezado(str(crudo) if crudo is not None else None, sedes, extra)
            for fila in range(bloque.fila_inicio, bloque.fila_fin):
                slot = _leer_slot_clasificado(
                    hoja,
                    fila,
                    columna,
                    bloque.col_inicio,
                    bloque.fecha,
                    terapeuta,
                    sede,
                    "horizontal",
                    archivo,
                    bloque.hoja,
                    incluir_vacios,
                    hora_invalidas,
                )
                if slot is not None:
                    slots.append(slot)
    if errores:
        raise ErrorDeModelo(errores)
    return slots, hora_invalidas


def ejecutar_clasificacion(
    libro: Workbook,
    sedes: tuple[str, ...],
    bloques: tuple | list,
    archivo: str,
    extra: dict[str, str] | None = None,
    layout: str = "vertical",
    incluir_vacios: bool = False,
    cerradas: bool = False,
) -> ResultadoClasificado:
    hora_invalidas: list[dict[str, str]] = []
    if layout == "vertical":
        slots, hora_invalidas = extraer_clasificado_vertical(
            libro, sedes, bloques, archivo, extra, incluir_vacios
        )
    elif layout == "horizontal":
        slots, hora_invalidas = extraer_clasificado_horizontal(
            libro, sedes, bloques, archivo, extra, incluir_vacios
        )
    else:
        raise ErrorDeModelo(f"layout desconocido: {layout}")
    filas_internas = [_slot_a_fila_interna(s) for s in slots]
    filas_publicas = exportar_slots_publicos(slots, cerradas=cerradas)
    dias = tuple(sorted({b.fecha for b in bloques}))
    return ResultadoClasificado(
        slots=slots,
        filas_internas=filas_internas,
        filas_publicas=filas_publicas,
        filas_hora_invalida=hora_invalidas,
        dias=dias,
        archivo=archivo,
    )


def escribir_clasificacion_xlsx(
    ruta: Path,
    filas: list[dict[str, str]],
    columnas: list[str] | None = None,
) -> None:
    from openpyxl import Workbook as LibroNuevo

    columnas = columnas or COLUMNAS_SLOTS_INTERNAS
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro = LibroNuevo()
    hoja = libro.active
    hoja.title = "clasificacion"
    hoja.append(columnas)
    for fila in filas:
        hoja.append([fila.get(col, "") for col in columnas])
    libro.save(ruta)
    libro.close()


def extraer_vertical(
    libro: Workbook,
    sedes: tuple[str, ...],
    bloques: tuple[BloqueVertical, ...] | list[BloqueVertical],
    extra: dict[str, str] | None = None,
    excluir_interno: bool = True,
) -> list[SesionAgenda]:
    sesiones: list[SesionAgenda] = []
    errores: list[str] = []
    for bloque in bloques:
        if bloque.hoja not in libro.sheetnames:
            errores.append(f"no está la hoja {bloque.hoja}")
            continue
        hoja = libro[bloque.hoja]
        for columna in range(2, hoja.max_column + 1):
            crudo = hoja.cell(row=bloque.fila_encabezado, column=columna).value
            terapeuta, sede = partir_encabezado(str(crudo) if crudo is not None else None, sedes, extra)
            letra = get_column_letter(columna)
            for fila in range(bloque.fila_inicio, bloque.fila_fin):
                donde = f"{bloque.hoja} {bloque.fecha.isoformat()} columna {letra}"
                sesion = _leer_slot(
                    hoja, fila, columna, 1, bloque.fecha, terapeuta, sede,
                    "vertical", excluir_interno, donde, errores,
                )
                if sesion is not None:
                    sesiones.append(sesion)
    if errores:
        raise ErrorDeModelo(errores)
    return sesiones


def extraer_horizontal(
    libro: Workbook,
    sedes: tuple[str, ...],
    bloques: tuple[BloqueHorizontal, ...] | list[BloqueHorizontal],
    extra: dict[str, str] | None = None,
    excluir_interno: bool = True,
) -> list[SesionAgenda]:
    sesiones: list[SesionAgenda] = []
    errores: list[str] = []
    for bloque in bloques:
        if bloque.hoja not in libro.sheetnames:
            errores.append(f"no está la hoja {bloque.hoja}")
            continue
        hoja = libro[bloque.hoja]
        for columna in range(bloque.col_inicio + 1, bloque.col_fin):
            crudo = hoja.cell(row=bloque.fila_encabezado, column=columna).value
            terapeuta, sede = partir_encabezado(str(crudo) if crudo is not None else None, sedes, extra)
            letra = get_column_letter(columna)
            for fila in range(bloque.fila_inicio, bloque.fila_fin):
                donde = f"{bloque.hoja} {bloque.fecha.isoformat()} columna {letra}"
                sesion = _leer_slot(
                    hoja, fila, columna, bloque.col_inicio, bloque.fecha, terapeuta, sede,
                    "horizontal", excluir_interno, donde, errores,
                )
                if sesion is not None:
                    sesiones.append(sesion)
    if errores:
        raise ErrorDeModelo(errores)
    return sesiones


def exportar_sesiones_publicas(
    sesiones: list[SesionAgenda],
    cerradas: bool = False,
) -> list[dict[str, str | int]]:
    """Aliases, fecha, hora, folio e iu_if. Sin texto, terapeuta ni sede en claro."""
    terapeutas: dict[str, str] = {}
    sedes_cod: dict[str, str] = {}
    filas: list[dict[str, str | int]] = []
    for indice, sesion in enumerate(sesiones, start=1):
        fila: dict[str, str | int] = {
            "contacto_id": indice,
            "fecha": sesion.fecha.isoformat(),
            "terapeuta_alias": codigo(terapeutas, sesion.terapeuta, "T", 3),
            "sede_alias": codigo(sedes_cod, sesion.sede or "", "S", 2),
            "hora": sesion.hora.strftime("%H:%M:%S"),
        }
        if not cerradas:
            fila["folio"] = sesion.folio or ""
            fila["iu_if"] = sesion.iu_if or ""
            fila["layout"] = sesion.layout
        filas.append(fila)
    return filas


def ejecutar(
    libro: Workbook,
    sedes: tuple[str, ...],
    bloques: tuple | list,
    extra: dict[str, str] | None = None,
    layout: str = "vertical",
    excluir_interno: bool = True,
    cerradas: bool = False,
) -> Resultado:
    if layout == "vertical":
        sesiones = extraer_vertical(libro, sedes, bloques, extra, excluir_interno)
    elif layout == "horizontal":
        sesiones = extraer_horizontal(libro, sedes, bloques, extra, excluir_interno)
    else:
        raise ErrorDeModelo(f"layout desconocido: {layout}")
    if not sesiones:
        raise ErrorDeModelo("No hubo sesiones con marca de asistencia y texto.")
    filas = exportar_sesiones_publicas(sesiones, cerradas=cerradas)
    dias = tuple(sorted({b.fecha for b in bloques}))
    return Resultado(sesiones=sesiones, filas=filas, dias=dias)


def escribir_csv(ruta: Path, filas: list[dict[str, str | int]], columnas: list[str] | None = None) -> None:
    import csv

    columnas = columnas or list(filas[0].keys())
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with ruta.open("w", newline="", encoding="utf-8") as archivo:
        escritor = csv.DictWriter(archivo, fieldnames=columnas, lineterminator="\r\n")
        escritor.writeheader()
        escritor.writerows(filas)


def procesar(
    ruta_libro: Path,
    ruta_sedes: Path,
    bloques: tuple | list,
    layout: str,
    ruta_variantes: Path | None = None,
    excluir_interno: bool = True,
    cerradas: bool = False,
) -> Resultado:
    import openpyxl

    from sesiones.reglas import cargar_sedes

    if not ruta_libro.is_file():
        raise ErrorDeModelo(f"No está el libro: {ruta_libro}")
    sedes, extra = cargar_sedes(ruta_sedes, ruta_variantes)
    libro = openpyxl.load_workbook(ruta_libro, data_only=False)
    try:
        return ejecutar(libro, sedes, bloques, extra, layout, excluir_interno, cerradas)
    finally:
        libro.close()


def procesar_clasificacion(
    ruta_libro: Path,
    ruta_sedes: Path,
    bloques: tuple | list,
    layout: str,
    ruta_variantes: Path | None = None,
    incluir_vacios: bool = False,
    cerradas: bool = False,
) -> ResultadoClasificado:
    import openpyxl

    from sesiones.reglas import cargar_sedes

    if not ruta_libro.is_file():
        raise ErrorDeModelo(f"No está el libro: {ruta_libro}")
    sedes, extra = cargar_sedes(ruta_sedes, ruta_variantes)
    libro = openpyxl.load_workbook(ruta_libro, data_only=False)
    try:
        return ejecutar_clasificacion(
            libro,
            sedes,
            bloques,
            ruta_libro.name,
            extra,
            layout,
            incluir_vacios,
            cerradas,
        )
    finally:
        libro.close()
