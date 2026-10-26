"""Lectores de las bases de expedientes. Índices de columna fijos por layout (2025 / 2026)."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from sesiones.reglas import ErrorDeModelo, extraer_folio, extraer_iu_if, normalizar_sede, plegar

DEFAULT_HOJA_2025 = "Sheet1"
DEFAULT_HOJA_2026 = "Data"

# Layout export 2025 (columnas 1-based en openpyxl)
COL_2025 = {
    "terapeuta": 2,
    "sede": 3,
    "fecha": 4,
    "tipo": 6,
    "folio": 7,
    "estatus": 8,
    "sexo": 9,
    "rango": 10,
    "sustancia": 23,
    "numero_sesion": 26,
}

# Layout export 2026 (columnas 1-based en openpyxl)
COL_2026 = {
    "folio": 3,
    "terapeuta": 7,
    "sede": 8,
    "estatus_atencion": 9,
    "edad": 13,
    "genero": 14,
    "sustancia": 20,
    "fecha_primera": 28,
    "fecha_ultima": 29,
    "sesiones": {1: 30, 2: 31, 3: 32, 4: 33, 5: 34, 6: 35, 7: 36, 8: 37, 9: 38, 10: 39, 11: 40, 12: 41},
    "total": 43,
    "estatus_tratamiento": 44,
    "motivo_egreso": 46,
}


@dataclass(frozen=True)
class Evento2025:
    evento_id: int
    fecha: date | None
    terapeuta: str | None
    sede: str | None
    tipo: str | None
    folio: str | None
    estatus: str | None
    sexo: str | None
    rango_edad: str | None
    sustancia: str | None
    numero_sesion: str | None
    es_sesion: bool


@dataclass(frozen=True)
class Persona2026:
    folio: str
    terapeuta: str | None
    sede: str | None
    edad: str | None
    genero: str | None
    sustancia_impacto: str | None
    fecha_primera: date | None
    fecha_ultima: date | None
    sesiones_mes: dict[int, int]
    total_sesiones: int
    estatus_tratamiento: str | None
    motivo_egreso: str | None
    iu_if: str | None


def _texto(valor) -> str | None:
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _fecha(valor) -> date | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(texto[:19], fmt).date()
        except ValueError:
            continue
    return None


def _entero(valor) -> int:
    if valor is None or valor == "":
        return 0
    try:
        return int(float(valor))
    except (TypeError, ValueError):
        return 0


def mapear_tipo(texto: str | None) -> str | None:
    if not texto:
        return None
    t = plegar(texto)
    if "CONSEJER" in t:
        return "consejería"
    if "FAMILIAR" in t:
        return "familiar"
    if "USUARIO" in t:
        return "usuario"
    return None


def mapear_estatus(texto: str | None) -> str | None:
    if not texto:
        return None
    t = plegar(texto)
    if "ALTA POR TERMINO" in t or "ALTA POR TÉRMINO" in t:
        return "alta por término"
    if "ALTA POR INICIO" in t:
        return "alta por inicio"
    if "ASESOR" in t:
        return "asesoría"
    if "SESION" in t or "SEGUIMIENTO" in t:
        return "sesión"
    if "BAJA" in t or "ABANDONO" in t:
        return "baja"
    if "CANALIZ" in t:
        return "canalización"
    return None


def mapear_sexo(texto: str | None) -> str | None:
    if not texto:
        return None
    t = plegar(texto)
    if t.startswith("M") and "HOMBRE" not in t:
        return "M"
    if t.startswith("H"):
        return "H"
    return "NB"


def _es_sesion_evento(estatus: str | None) -> bool:
    return estatus == "sesión"


def leer_expediente_2025(libro, hoja: str = DEFAULT_HOJA_2025) -> list[Evento2025]:
    if hoja not in libro.sheetnames:
        raise ErrorDeModelo(f"no está la hoja {hoja}")
    ws = libro[hoja]
    eventos: list[Evento2025] = []
    for fila in range(2, ws.max_row + 1):
        def celda(clave: str):
            return ws.cell(row=fila, column=COL_2025[clave]).value

        if all(celda(k) is None or str(celda(k)).strip() == "" for k in ("fecha", "folio", "tipo", "estatus")):
            continue
        folio = extraer_folio(celda("folio"))
        estatus = mapear_estatus(_texto(celda("estatus")))
        eventos.append(
            Evento2025(
                evento_id=len(eventos) + 1,
                fecha=_fecha(celda("fecha")),
                terapeuta=_texto(celda("terapeuta")),
                sede=normalizar_sede(_texto(celda("sede"))),
                tipo=mapear_tipo(_texto(celda("tipo"))),
                folio=folio,
                estatus=estatus,
                sexo=mapear_sexo(_texto(celda("sexo"))),
                rango_edad=_texto(celda("rango")),
                sustancia=_texto(celda("sustancia")),
                numero_sesion=_texto(celda("numero_sesion")),
                es_sesion=_es_sesion_evento(estatus),
            )
        )
    return eventos


def leer_expediente_2026(libro, hoja: str = DEFAULT_HOJA_2026) -> list[Persona2026]:
    if hoja not in libro.sheetnames:
        raise ErrorDeModelo(f"no está la hoja {hoja}")
    ws = libro[hoja]
    personas: list[Persona2026] = []
    for fila in range(2, ws.max_row + 1):
        folio = extraer_folio(ws.cell(row=fila, column=COL_2026["folio"]).value)
        if folio is None:
            continue
        sesiones = {
            mes: _entero(ws.cell(row=fila, column=col).value)
            for mes, col in COL_2026["sesiones"].items()
        }
        total = _entero(ws.cell(row=fila, column=COL_2026["total"]).value)
        if total == 0:
            total = sum(sesiones.values())
        personas.append(
            Persona2026(
                folio=folio,
                terapeuta=_texto(ws.cell(row=fila, column=COL_2026["terapeuta"]).value),
                sede=normalizar_sede(_texto(ws.cell(row=fila, column=COL_2026["sede"]).value)),
                edad=_texto(ws.cell(row=fila, column=COL_2026["edad"]).value),
                genero=_texto(ws.cell(row=fila, column=COL_2026["genero"]).value),
                sustancia_impacto=_texto(ws.cell(row=fila, column=COL_2026["sustancia"]).value),
                fecha_primera=_fecha(ws.cell(row=fila, column=COL_2026["fecha_primera"]).value),
                fecha_ultima=_fecha(ws.cell(row=fila, column=COL_2026["fecha_ultima"]).value),
                sesiones_mes=sesiones,
                total_sesiones=total,
                estatus_tratamiento=_texto(ws.cell(row=fila, column=COL_2026["estatus_tratamiento"]).value),
                motivo_egreso=_texto(ws.cell(row=fila, column=COL_2026["motivo_egreso"]).value),
                iu_if=extraer_iu_if(folio),
            )
        )
    return personas


def cargar_libro(ruta: Path, data_only: bool = True):
    import openpyxl

    if not ruta.is_file():
        raise ErrorDeModelo(f"No está el libro: {ruta}")
    return openpyxl.load_workbook(ruta, data_only=data_only)
